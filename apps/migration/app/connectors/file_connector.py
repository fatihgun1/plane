"""
app/connectors/file_connector.py
==================================
CSV / XLSX import connector for the NTT-PMO Optimization Tool.

Reads a flat spreadsheet of work items and returns Item dataclasses.
Supports both .csv and .xlsx/.xlsm formats.

Expected columns (case-insensitive, extras ignored)
-----------------------------------------------------
    jira_key        | Required | e.g. "JUM-123"
    summary         | Required | free text
    issue_type      | Optional | defaults to "WRICEF"
    module          | Optional |
    stream          | Optional |
    priority        | Optional | int 1-5 or text (Highest/High/Medium/Low/Lowest)
    effort_days     | Optional | numeric
    status          | Optional |
    epic_key        | Optional |

Usage
-----
    from app.connectors.file_connector import FileConnector
    fc = FileConnector()
    items = fc.load("path/to/items.xlsx")
"""

from __future__ import annotations

import csv
import logging
import os
from datetime import datetime
from typing import Optional

from app.models import Item

logger = logging.getLogger(__name__)

_PRIORITY_MAP = {
    "highest": 1,
    "high": 2,
    "medium": 3,
    "low": 4,
    "lowest": 5,
}

# Jumbo xlsm uses "N - Label" prefixed strings for dev types
_ISSUE_TYPE_NORMALIZE: dict[str, str] = {
    "enhancement": "Enhancement",
    "data migration": "Data Migration",
    "interface": "Interface",
    "report": "Report",
    "conversion": "Conversion",
    "form": "Form",
    "workflow": "Workflow",
    "configuration": "Configuration",
}

# Canonical column names mapped from common aliases
# Jumbo xlsm-specific aliases added alongside generic ones
_COL_ALIASES: dict[str, list[str]] = {
    "jira_key":   ["jira_key", "key", "issue key", "jira key", "id"],
    "summary":    ["summary", "title", "name"],
    "issue_type": ["issue_type", "issuetype", "type", "issue type",
                   "development type"],
    "module":     ["module", "area", "responsible module"],
    "stream":     ["stream", "workstream", "work stream"],
    "priority":   ["priority"],
    "effort_days":["effort_days", "effort days", "effort", "story points",
                   "storypoints", "sp", "days", "total effort", "module effort"],
    "status":     ["status", "state"],
    "epic_key":   ["epic_key", "epic key", "epic", "epic link"],
    "description":["description"],
}


def _build_col_map(header: list[str]) -> dict[str, int]:
    """
    Map canonical field names to column indices by checking aliases.
    Returns only the fields that were found.
    """
    normalised = [h.strip().lower() for h in header]
    col_map: dict[str, int] = {}
    for canonical, aliases in _COL_ALIASES.items():
        for alias in aliases:
            if alias in normalised:
                col_map[canonical] = normalised.index(alias)
                break
    return col_map


def _coerce_priority(raw: str) -> Optional[int]:
    """Convert priority cell (int string or label) to 1-5 int.

    Handles formats:
      - plain int "2"
      - plain label "High"
      - Jumbo "N - Label" e.g. "1 - High", "2 - Medium"
    """
    raw = raw.strip()
    if not raw:
        return None
    # Jumbo format: "1 - High" → take the label part
    if " - " in raw:
        parts = raw.split(" - ", 1)
        # Try the label part first
        label = parts[1].strip().lower()
        mapped = _PRIORITY_MAP.get(label)
        if mapped:
            return mapped
        # Fall back to the numeric prefix
        if parts[0].strip().isdigit():
            val = int(parts[0].strip())
            return val if 1 <= val <= 5 else None
    if raw.isdigit():
        val = int(raw)
        return val if 1 <= val <= 5 else None
    return _PRIORITY_MAP.get(raw.lower())


def _normalize_issue_type(raw: str) -> str:
    """Normalise issue_type / development type.

    Strips leading "N - " prefix from Jumbo format and maps to canonical label.
    Falls back to the original value if not in map.
    """
    raw = raw.strip()
    label = raw
    if " - " in raw:
        label = raw.split(" - ", 1)[1].strip()
    return _ISSUE_TYPE_NORMALIZE.get(label.lower(), label) if label else raw


def _coerce_float(raw: str) -> Optional[float]:
    try:
        return float(raw.strip()) if raw.strip() else None
    except ValueError:
        return None


def _row_to_item(
    row: dict[str, str],
    col_map: dict[str, int],
    row_values: list[str],
    imported_at: str,
) -> Optional[Item]:
    """Convert a raw row (list of cell strings) to an Item. Returns None if jira_key missing."""
    def get(field: str) -> str:
        idx = col_map.get(field)
        if idx is None or idx >= len(row_values):
            return ""
        return str(row_values[idx]).strip()

    jira_key = get("jira_key")
    if not jira_key:
        return None

    priority_raw = get("priority")
    priority = _coerce_priority(priority_raw) if priority_raw else None

    effort_raw = get("effort_days")
    effort_days = _coerce_float(effort_raw)

    issue_type_raw = get("issue_type")
    issue_type = _normalize_issue_type(issue_type_raw) if issue_type_raw else "WRICEF"

    return Item(
        id=None,
        jira_key=jira_key,
        summary=get("summary") or None,
        issue_type=issue_type,
        module=get("module") or None,
        stream=get("stream") or None,
        priority=priority,
        effort_days=effort_days,
        status=get("status") or None,
        epic_key=get("epic_key") or None,
        imported_at=imported_at,
    )


class FileConnector:
    """
    Loads work items from CSV or Excel files.

    Parameters
    ----------
    encoding : str
        Encoding used when reading CSV files (default: utf-8-sig).
    """

    def __init__(self, encoding: str = "utf-8-sig") -> None:
        self.encoding = encoding

    # ------------------------------------------------------------------
    # Public entry point
    # ------------------------------------------------------------------

    def load(self, path: str, sheet_name: Optional[str] = None) -> list[Item]:
        """
        Load items from *path* (.csv, .xlsx, or .xlsm).

        Parameters
        ----------
        path : str
            Path to the file.
        sheet_name : str, optional
            Excel sheet to read. If None uses the active sheet.
            Ignored for CSV files.

        Returns a list of Item objects. Rows missing jira_key are skipped.
        """
        ext = os.path.splitext(path)[1].lower()
        if ext == ".csv":
            return self._load_csv(path)
        if ext in (".xlsx", ".xlsm", ".xls"):
            return self._load_excel(path, sheet_name=sheet_name)
        raise ValueError(f"Unsupported file type: {ext!r}")

    # ------------------------------------------------------------------
    # CSV loader
    # ------------------------------------------------------------------

    def _load_csv(self, path: str) -> list[Item]:
        imported_at = datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%S")
        items: list[Item] = []

        with open(path, newline="", encoding=self.encoding) as fh:
            reader = csv.reader(fh)
            header = next(reader, None)
            if header is None:
                logger.warning("CSV file is empty: %s", path)
                return []
            col_map = _build_col_map(header)
            if "jira_key" not in col_map:
                raise ValueError(
                    f"CSV {path!r} has no recognised jira_key column. "
                    f"Header: {header}"
                )
            for line_no, row_values in enumerate(reader, start=2):
                item = _row_to_item({}, col_map, row_values, imported_at)
                if item:
                    items.append(item)
                else:
                    logger.debug("Skipped row %d (no jira_key)", line_no)

        logger.info("FileConnector CSV: loaded %d items from %s", len(items), path)
        return items

    # ------------------------------------------------------------------
    # Excel loader (openpyxl)
    # ------------------------------------------------------------------

    def _load_excel(self, path: str, sheet_name: Optional[str] = None) -> list[Item]:
        try:
            import openpyxl  # noqa: PLC0415
        except ImportError as exc:
            raise ImportError(
                "openpyxl is required to read Excel files. "
                "Install it with: pip install openpyxl"
            ) from exc

        imported_at = datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%S")
        items: list[Item] = []

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if sheet_name:
            if sheet_name not in wb.sheetnames:
                wb.close()
                raise ValueError(
                    f"Sheet {sheet_name!r} not found. "
                    f"Available: {wb.sheetnames}"
                )
            ws = wb[sheet_name]
        else:
            ws = wb.active

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            logger.warning("Excel file has no rows: %s", path)
            wb.close()
            return []

        header = [str(c) if c is not None else "" for c in rows[0]]
        col_map = _build_col_map(header)
        if "jira_key" not in col_map:
            wb.close()
            raise ValueError(
                f"Excel {path!r} has no recognised jira_key column. "
                f"Header: {header}"
            )

        for line_no, raw_row in enumerate(rows[1:], start=2):
            row_values = [str(c) if c is not None else "" for c in raw_row]
            item = _row_to_item({}, col_map, row_values, imported_at)
            if item:
                items.append(item)
            else:
                logger.debug("Skipped row %d (no jira_key)", line_no)

        wb.close()
        logger.info("FileConnector Excel: loaded %d items from %s", len(items), path)
        return items

    # ------------------------------------------------------------------
    # Convenience: detect sheet names (Excel only)
    # ------------------------------------------------------------------

    def list_sheets(self, path: str) -> list[str]:
        """Return sheet names for an Excel file."""
        try:
            import openpyxl  # noqa: PLC0415
        except ImportError as exc:
            raise ImportError("openpyxl required") from exc
        wb = openpyxl.load_workbook(path, read_only=True)
        names = wb.sheetnames
        wb.close()
        return names
